# -*- coding: utf-8 -*-
"""
简历批量投递助手 - 本地服务
功能：投递清单管理（增删改查 / Excel 导入导出）、邮件话术（可存好几套，按「投递定位」取用）、
     证明材料合并成一个 PDF（PDF + 图片混拼，图片自动摆正 / 压到 A4）、
     QQ 邮箱 SMTP 批量发送（间隔控制 / 每日上限 / 失败记录 / 中途停止）
说明：所有数据只保存在本机 data/ 与 attachments/ 目录中，不经过任何第三方服务器。
"""
import base64
import io
import json
import os
import re
import shutil
import smtplib
import socket
import subprocess
import sys
import tempfile
import threading
import time
import webbrowser
from datetime import datetime
from email.header import Header
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.utils import formataddr
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, quote, urlparse

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE, "data")
ATTACH_DIR = os.path.join(BASE, "attachments")
INDEX_HTML = os.path.join(BASE, "index.html")

ATT_NOTE_FILE = "attachment_notes.json"   # {文件名: 备注文字}，备注只存在本机，不会随邮件发出
SHEET_FILE = "sheet_draft.json"           # 在线填表的草稿，关掉软件再打开还在

# 模板库：可以同时存好几套话术（一套投 AI 岗、一套投专业对口岗…）。
# 每条投递目标在「投递定位」那一列挑一套；不挑就跟随这里标记为默认的那套。
TPL_FILE = "templates.json"
DEFAULT_TPL_NAME = "默认"

# 在线填表 / 按列导入 的列顺序：第1列公司、第2列岗位、第3列收件邮箱、第4列备注
SHEET_COLS = ["company", "position", "email", "note"]
SHEET_COLS_CN = ["公司", "岗位", "收件邮箱", "备注"]
MAX_TAGS_PER_JOB = 12                     # 一条最多几个标签，防止乱填撑爆界面

# 两封邮件之间最少隔多少秒。这不是性能问题，是账号安全问题：
# 邮箱服务商看不懂"你在投简历"，它只看"一个账号短时间内发给一堆互不相识的人、内容还差不多"，
# 这就是垃圾邮件的典型画像，频率越高越先被拦。所以再想快也不让填到 20 秒以下。
MIN_SEND_INTERVAL = 20
SAFE_SEND_INTERVAL = 30                   # 界面上低于这个数会提示"太快了"

# ---------------------------------------------------------------- 材料合并
# 暂存区：用户丢进来的证明材料先放这儿，排好顺序再合成。
# 顺序直接写进文件名开头的序号（001__、002__…），所以关掉软件再打开顺序也不会乱。
MERGE_DIR = os.path.join(DATA_DIR, "_merge_staging")

A4_W, A4_H = 595.28, 841.89   # A4 纸尺寸，单位「点」（1 点 = 1/72 英寸）
MERGE_DPI = 144               # 图片放进 PDF 的分辨率。144 够 HR 看清，文件也不会太大
MERGE_MARGIN = 0.04           # 页面四周留 4% 白边，证书照片贴着纸边不好看、打印也容易被切

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".gif", ".tif", ".tiff"}

# 能在网页里直接打开预览的格式：PDF 用浏览器自带的阅读器（能放大、翻页、搜索），图片直接显示。
# 其它格式（zip、doc 之类）浏览器打不开，界面里会提示改用下载。
PREVIEW_TYPES = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
    ".webp": "image/webp",
}

SMTP_HOST = "smtp.qq.com"
SMTP_PORT = 465

DEFAULT_CONFIG = {
    "email": "",          # QQ 邮箱地址
    "auth_code": "",      # QQ 邮箱授权码（不是登录密码）
    "display_name": "",   # 发件人显示名
    "name": "",           # 我的姓名（模板变量 {姓名}）
    "school": "",         # 我的学校（模板变量 {学校}）
    "major": "",          # 我的专业（模板变量 {专业}）
    "interval": 40,       # 每封邮件间隔（秒）
    "daily_limit": 50,    # 每日发送上限
}

DEFAULT_TEMPLATE = {
    "subject": "应聘【{岗位}】- {姓名}（{学校}）",
    "body": (
        "尊敬的HR：\n\n"
        "您好！\n\n"
        "我是{学校}{姓名}，从招聘信息中了解到贵司正在招聘{岗位}，"
        "我对该岗位非常感兴趣，现将我的简历附上，期待能有机会参加面试。\n\n"
        "感谢您在百忙之中查阅我的邮件，期待您的回复！\n\n"
        "祝工作顺利！\n\n"
        "{姓名}\n"
        "{学校}\n"
    ),
}


# ---------------------------------------------------------------- 数据存取
# 一把全局锁：所有 json 的"读出来 → 改 → 写回去"都必须整段在锁里完成。
# 不然两个操作同时发生（发信线程 + 网页上的改动），后写的那次会把先写的覆盖掉。
DATA_LOCK = threading.RLock()
# 读坏 / 写坏过的文件，会在网页顶部提醒用户
DATA_WARNINGS = []


def _path(name):
    return os.path.join(DATA_DIR, name)


def _warn(text):
    if text not in DATA_WARNINGS:
        DATA_WARNINGS.append(text)
    print("[警告] %s" % text)


def _quarantine(path, why):
    """文件内容读不出来时，先改名留一份，绝不直接覆盖 —— 原文件是用户唯一的底稿。"""
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = "%s.corrupt-%s" % (path, stamp)
    try:
        os.replace(path, backup)
    except OSError:
        backup = path
    _warn("%s 的内容读不出来（%s），已把原文件备份成「%s」，软件先按「没有数据」显示。"
          % (os.path.basename(path), why, os.path.basename(backup)))
    return backup


def load_json(name, default):
    p = _path(name)
    with DATA_LOCK:
        if not os.path.exists(p):
            return default
        try:
            f = open(p, "r", encoding="utf-8")
        except FileNotFoundError:
            return default
        except OSError as e:
            # 文件只是暂时打不开（被别的程序占着之类），这时绝不能动它
            _warn("%s 暂时读不出来（%s），本次按空数据处理，原文件没有改动。"
                  % (os.path.basename(p), e))
            return default
        try:
            with f:
                return json.load(f)
        except Exception as e:
            _quarantine(p, e)
            return default


def save_json(name, obj):
    """"先写临时文件、再改名"。临时文件名带进程号和线程号 —— 两个请求同时保存也不会互相踩。"""
    with DATA_LOCK:
        p = _path(name)
        tmp = "%s.%d.%d.tmp" % (p, os.getpid(), threading.get_ident())
        try:
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(obj, f, ensure_ascii=False, indent=2)
            os.replace(tmp, p)
        finally:
            if os.path.exists(tmp):
                try:
                    os.remove(tmp)
                except OSError:
                    pass


def init_dirs():
    os.makedirs(DATA_DIR, exist_ok=True)
    os.makedirs(ATTACH_DIR, exist_ok=True)
    os.makedirs(MERGE_DIR, exist_ok=True)
    if not os.path.exists(_path("config.json")):
        save_json("config.json", DEFAULT_CONFIG)
    if not os.path.exists(_path("jobs.json")):
        save_json("jobs.json", [])
    if not os.path.exists(_path(TPL_FILE)):
        get_templates()      # 第一次运行自动建一套「默认」；老版本升上来的会把 template.json 的内容搬进去
    if not os.path.exists(_path("sent_log.json")):
        save_json("sent_log.json", [])
    if not os.path.exists(_path(ATT_NOTE_FILE)):
        save_json(ATT_NOTE_FILE, {})
    if not os.path.exists(_path(SHEET_FILE)):
        save_json(SHEET_FILE, {"rows": [], "updated_at": ""})


def get_config():
    cfg = dict(DEFAULT_CONFIG)
    cfg.update(load_json("config.json", {}))
    return cfg


# ---------------------------------------------------------------- 模板库
def _new_tpl_id():
    return "t%d" % int(time.time() * 1000)


def _clean_tpl(t):
    """把一套模板收拾成固定格式，坏数据一律换成安全值。"""
    return {
        "id": str(t.get("id") or _new_tpl_id()),
        "name": (str(t.get("name") or "").strip() or DEFAULT_TPL_NAME)[:30],
        "subject": str(t.get("subject") or ""),
        "body": str(t.get("body") or ""),
    }


def get_templates():
    """读模板库。

    老版本只有一套（data/template.json），第一次读不到 templates.json 就自动搬过来，
    所以用户升级后不用重新写话术。
    """
    d = load_json(TPL_FILE, None)
    items = d.get("list") if isinstance(d, dict) else None
    if not isinstance(items, list) or not items:
        old = load_json("template.json", {})
        old = old if isinstance(old, dict) else {}
        base = dict(DEFAULT_TEMPLATE)
        base.update(old)
        d = {"list": [{"id": "t1", "name": DEFAULT_TPL_NAME,
                       "subject": base.get("subject", ""), "body": base.get("body", "")}],
             "default_id": "t1"}
        save_json(TPL_FILE, d)

    clean = [_clean_tpl(t) for t in d.get("list", []) if isinstance(t, dict)]
    if not clean:
        clean = [{"id": "t1", "name": DEFAULT_TPL_NAME,
                  "subject": DEFAULT_TEMPLATE["subject"], "body": DEFAULT_TEMPLATE["body"]}]
    default_id = str(d.get("default_id") or "")
    if default_id not in [t["id"] for t in clean]:
        default_id = clean[0]["id"]
    return {"list": clean, "default_id": default_id}


def find_template(tpls, tid):
    """按 id 找一套；找不到（比如那条清单指定的模板被删了）就回落到默认那套。"""
    for t in tpls["list"]:
        if t["id"] == tid:
            return t
    for t in tpls["list"]:
        if t["id"] == tpls["default_id"]:
            return t
    return tpls["list"][0]


def template_for_job(job, tpls):
    """这条投递目标该用哪套话术：它自己指定的 > 默认那套。"""
    return find_template(tpls, str((job or {}).get("template_id") or ""))


def tpl_display_name(job, tpls):
    """这条清单的「投递定位」显示成什么：没指定就写「默认」，指定了写那套的名字。

    导出 Excel、网页上都要用同一个写法，免得两边对不上。
    """
    tid = str((job or {}).get("template_id") or "")
    if tid:
        for t in tpls["list"]:
            if t["id"] == tid:
                return t["name"]
    return DEFAULT_TPL_NAME


def save_templates(tpls):
    save_json(TPL_FILE, tpls)
    return tpls


def get_jobs():
    jobs = load_json("jobs.json", [])
    return jobs if isinstance(jobs, list) else []


def save_jobs(jobs):
    save_json("jobs.json", jobs)


def mutate_jobs(fn):
    """清单的"读最新 → 改 → 写回"必须整段锁在一起。

    以前是"先把整份清单读进内存，改完再整份写回去"，
    结果发信线程手上的旧副本会把用户在网页上的改动全部盖掉（删掉的公司还会复活）。
    凡是要改清单的地方都走这个函数。
    """
    with DATA_LOCK:
        jobs = get_jobs()
        fn(jobs)
        save_jobs(jobs)
        return jobs


def mutate_job(jid, fn):
    """只动清单里的某一条，其他条目保持磁盘上的最新状态。没找到返回 False。"""
    hit = []

    def _run(jobs):
        for j in jobs:
            if str(j.get("id")) == str(jid):
                fn(j)
                hit.append(True)
                break

    mutate_jobs(_run)
    return bool(hit)


def find_job(jid):
    for j in get_jobs():
        if str(j.get("id")) == str(jid):
            return j
    return None


def jobs_stamp():
    """清单文件的改动时间。网页轮询时拿它比较，只有真变了才重画表格
    （不然每 2.5 秒重画一次，会把你正在填的那个格子打断）。"""
    try:
        return os.path.getmtime(_path("jobs.json"))
    except OSError:
        return 0


def today_sent_count():
    today = datetime.now().strftime("%Y-%m-%d")
    logs = load_json("sent_log.json", [])
    if not isinstance(logs, list):
        return 0
    return sum(1 for e in logs
               if isinstance(e, dict) and str(e.get("sent_at", "")).startswith(today) and e.get("ok"))


def append_sent_log(entry):
    with DATA_LOCK:
        logs = load_json("sent_log.json", [])
        if not isinstance(logs, list):
            logs = []
        logs.append(entry)
        save_json("sent_log.json", logs)


# ---------------------------------------------------------------- 变量渲染
def render(tpl_text, job, cfg):
    mapping = {
        "{公司}": job.get("company", ""),
        "{岗位}": job.get("position", ""),
        "{备注}": job.get("note", ""),
        "{姓名}": cfg.get("name", ""),
        "{学校}": cfg.get("school", ""),
        "{专业}": cfg.get("major", ""),
    }
    out = tpl_text or ""
    for k, v in mapping.items():
        out = out.replace(k, v)
    return out


# ---------------------------------------------------------------- 标签
def split_tags(v, limit=None):
    """把标签统一成去重后的列表，顺便告诉你被挤掉几个。

    前端传列表 ["9/19投", "国企"] 或字符串 "9/19投,国企" 都能收。
    逗号（中英文）、顿号、分号、竖线、空格都当分隔符 ——
    中文里"国企、内推、9/19投"很常见，以前顿号不算分隔符，会被当成一个超长标签。
    返回 (要存的标签, 被挤掉的个数) —— 第二条是给界面提示用的：
    以前超了就默默丢掉，用户以为"没显示"，其实是压根没存进去。
    """
    limit = MAX_TAGS_PER_JOB if limit is None else limit
    if isinstance(v, (list, tuple)):
        raw = " ".join(str(x) for x in v)
    else:
        raw = str(v or "")
    out = []
    for p in re.split(r"[,，、;；|｜\s]+", raw):
        p = p.strip()
        if p and p not in out:
            out.append(p)
    return out[:limit], max(0, len(out) - limit)


def norm_tags(v):
    """只要标签列表、不关心被挤掉几个的时候用这个。"""
    return split_tags(v)[0]


def tag_all():
    """清单里用过的所有标签，给筛选下拉框用。"""
    seen = []
    for j in get_jobs():
        for t in (j.get("tags") or []):
            if t and t not in seen:
                seen.append(t)
    return seen


# ---------------------------------------------------------------- 邮件发送
def get_att_notes():
    d = load_json(ATT_NOTE_FILE, {})
    return d if isinstance(d, dict) else {}


def set_att_note(name, note):
    """写入/清除某个附件的备注（备注为空等于删掉备注）。备注只存本机，不会写进邮件。"""
    d = get_att_notes()
    note = (note or "").strip()
    if note:
        d[name] = note
    else:
        d.pop(name, None)
    save_json(ATT_NOTE_FILE, d)
    return d


def drop_att_note(name):
    d = get_att_notes()
    if name in d:
        d.pop(name)
        save_json(ATT_NOTE_FILE, d)


def preview_kind(name):
    """"pdf" / "image" / "other" —— 决定浏览器里能不能直接打开看。"""
    ext = os.path.splitext(str(name or ""))[1].lower()
    if ext not in PREVIEW_TYPES:
        return "other"
    return "pdf" if ext == ".pdf" else "image"


def list_attachments():
    notes = get_att_notes()
    items = []
    for fn in sorted(os.listdir(ATTACH_DIR)):
        p = os.path.join(ATTACH_DIR, fn)
        if os.path.isfile(p):
            items.append({"name": fn, "size": os.path.getsize(p),
                          "note": notes.get(fn, ""), "kind": preview_kind(fn)})
    return items


# ---------------------------------------------------------------- 材料合并
def safe_fname(name):
    """只保留文件名本身（去掉路径），再把 Windows 不认的字符换成下划线。"""
    name = os.path.basename(str(name or "")).strip()
    for ch in '\\/:*?"<>|':
        name = name.replace(ch, "_")
    return name or "未命名"


def is_pdf_name(name):
    return os.path.splitext(name)[1].lower() == ".pdf"


def is_image_name(name):
    return os.path.splitext(name)[1].lower() in IMAGE_EXTS


def pdf_page_count(path):
    """读 PDF 页数；读不了（加密/损坏）返回 None。"""
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:
                pass
        return len(reader.pages)
    except Exception:
        return None


def _staged_files():
    """暂存区里所有文件，按文件名开头的序号排好。"""
    if not os.path.isdir(MERGE_DIR):
        return []
    out = []
    for fn in sorted(os.listdir(MERGE_DIR)):
        p = os.path.join(MERGE_DIR, fn)
        if os.path.isfile(p) and "__" in fn:
            out.append((fn, fn.split("__", 1)[1]))
    return out


def merge_list():
    items = []
    for stored, shown in _staged_files():
        p = os.path.join(MERGE_DIR, stored)
        kind = "pdf" if is_pdf_name(shown) else ("image" if is_image_name(shown) else "other")
        items.append({
            "id": stored,
            "name": shown,
            "size": os.path.getsize(p),
            "kind": kind,
            "ok": kind in ("pdf", "image"),
            "pages": pdf_page_count(p) if kind == "pdf" else (1 if kind == "image" else 0),
        })
    return items


def merge_add(name, data):
    """把上传的文件放进暂存区，序号续在最后。"""
    os.makedirs(MERGE_DIR, exist_ok=True)
    shown = safe_fname(name)
    used = [int(f.split("__", 1)[0]) for f, _ in _staged_files() if f.split("__", 1)[0].isdigit()]
    seq = (max(used) + 1) if used else 1
    while True:
        target = os.path.join(MERGE_DIR, "%03d__%s" % (seq, shown))
        if not os.path.exists(target):
            break
        seq += 1
    with open(target, "wb") as f:
        f.write(data)
    return target


def merge_move(fid, direction):
    """把某个文件往上/往下挪一格：交换两个文件的序号。"""
    ids = [stored for stored, _ in _staged_files()]
    if fid not in ids:
        return
    i = ids.index(fid)
    j = i + (-1 if direction < 0 else 1)
    if j < 0 or j >= len(ids):
        return
    a, b = ids[i], ids[j]
    seq_a, name_a = a.split("__", 1)
    seq_b, name_b = b.split("__", 1)
    tmp = os.path.join(MERGE_DIR, "000__swap.tmp")
    os.replace(os.path.join(MERGE_DIR, a), tmp)
    os.replace(os.path.join(MERGE_DIR, b), os.path.join(MERGE_DIR, seq_a + "__" + name_b))
    os.replace(tmp, os.path.join(MERGE_DIR, seq_b + "__" + name_a))


def merge_remove(fid):
    p = os.path.join(MERGE_DIR, os.path.basename(fid))
    if os.path.isfile(p):
        os.remove(p)
    # 删掉一个以后，把剩下的序号重新排成 001、002… 保持连续
    rest = _staged_files()
    for new_i, (stored, shown) in enumerate(rest, 1):
        want = "%03d__%s" % (new_i, shown)
        if want != stored:
            os.replace(os.path.join(MERGE_DIR, stored), os.path.join(MERGE_DIR, want))


def merge_clear():
    for stored, _ in _staged_files():
        try:
            os.remove(os.path.join(MERGE_DIR, stored))
        except OSError:
            pass


def render_image_page(path):
    """把一张图片画到一页 A4 白纸上：按比例缩到放得下，居中，四周留白。

    手机拍的照片常带"躺倒"的方向信息，先按它摆正。
    图是横的就用 A4 横向，是竖的就用 A4 纵向，免得横版证书被压得很小。
    返回 (PIL 图片, 输出 DPI)。
    """
    from PIL import Image, ImageOps

    im = Image.open(path)
    im.load()
    fixed = ImageOps.exif_transpose(im)
    if fixed is not None:
        im = fixed

    # 透明背景（PNG）先垫成白的，不然 PDF 里会变黑块
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, "white")
        bg.paste(im, mask=im.split()[-1])
        im = bg
    elif im.mode != "RGB":
        im = im.convert("RGB")

    pw, ph = (A4_H, A4_W) if im.width > im.height else (A4_W, A4_H)
    k = MERGE_DPI / 72.0
    cw = max(1, int(round(pw * k)))
    ch = max(1, int(round(ph * k)))
    canvas = Image.new("RGB", (cw, ch), "white")

    mx = int(cw * MERGE_MARGIN)
    my = int(ch * MERGE_MARGIN)
    box_w = max(1, cw - mx * 2)
    box_h = max(1, ch - my * 2)
    ratio = min(box_w / im.width, box_h / im.height)
    nw = max(1, int(round(im.width * ratio)))
    nh = max(1, int(round(im.height * ratio)))
    if (nw, nh) != im.size:
        im = im.resize((nw, nh), Image.LANCZOS)
    canvas.paste(im, ((cw - nw) // 2, (ch - nh) // 2))

    # 用页面宽度反推 DPI，让 PDF 里的页面尺寸正好等于 A4（避免取整误差）
    return canvas, cw * 72.0 / pw


_GUI_PY = {"exe": None}

LOCAL_PY_FILE = os.path.join(DATA_DIR, "local_python.txt")


def local_python_paths():
    """读 data/local_python.txt：一行一个 Python 路径（本机专用，已排除在 Git 之外）。

    用途：自动探测万一找不到合适的解释器，用户可以在这里手写一个。
    """
    out = []
    try:
        with open(LOCAL_PY_FILE, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#"):
                    out.append(line)
    except OSError:
        pass
    return out


def find_gui_python():
    """找一个同时装了 tkinter 和 pypdf 的 Python，用来跑独立合并工具。

    坑：本软件自己跑在精简环境里，那个 Python 只有 pypdf、没有 tkinter，
    所以不能直接用 sys.executable（会报 No module named 'tkinter'，窗口根本开不出来）。
    按顺序挨个试，探测结果缓存起来。一个都找不到时返回 None。
    """
    if _GUI_PY["exe"]:
        return _GUI_PY["exe"]

    here = sys.executable or ""
    home = os.environ.get("USERPROFILE") or os.path.expanduser("~")
    cands = []
    for c in (
        *local_python_paths(),
        os.path.join(os.path.dirname(here), "pythonw.exe") if here else "",
        here,
        shutil.which("pythonw") or "",
        shutil.which("python") or "",
        os.path.join(home, "miniconda3", "python.exe") if home else "",
        os.path.join(home, "anaconda3", "python.exe") if home else "",
        r"C:\ProgramData\miniconda3\python.exe",
        r"C:\ProgramData\Anaconda3\python.exe",
        "pythonw",
        "python",
    ):
        if c and c not in cands:
            cands.append(c)

    flags = 0x08000000 if os.name == "nt" else 0
    for exe in cands:
        try:
            r = subprocess.run([exe, "-c", "import tkinter, pypdf"],
                               capture_output=True, timeout=25, creationflags=flags)
        except Exception:
            continue
        if r.returncode == 0:
            _GUI_PY["exe"] = exe
            return exe
    return None


def open_merge_tool(out_name=""):
    """把独立的「PDF / 图片 合并工具」作为单独窗口打开。

    它功能比网页版全（预览每一页、单张旋转、逐张调页面大小），所以主推这个。
    通过 --outdir 告诉它默认存到附件库，合完的 PDF 直接就能勾选发送。
    """
    script = os.path.join(BASE, "pdf_merge_gui.py")
    if not os.path.isfile(script):
        raise ValueError("找不到独立合并工具 pdf_merge_gui.py，请确认它还在软件目录里")

    exe = find_gui_python()
    if not exe:
        raise ValueError("没找到能开窗口的 Python（需要同时装了 tkinter 和 pypdf）。"
                         "可以双击软件目录里的「合并PDF.bat」手动打开。")

    cmd = [exe, script, "--outdir", ATTACH_DIR]
    if out_name:
        cmd += ["--outname", safe_fname(out_name)]
    kwargs = {"cwd": BASE}
    if os.name == "nt":
        # 不额外弹一个黑色控制台窗口（工具自己的窗口照常显示）
        kwargs["creationflags"] = 0x08000000
    subprocess.Popen(cmd, **kwargs)


def merge_run(out_name):
    """按暂存区顺序合成一个 PDF，直接放进附件库。返回 (文件名, 页数, 字节数)。"""
    from pypdf import PdfReader, PdfWriter

    picked = [it for it in merge_list() if it["ok"]]
    if not picked:
        raise ValueError("暂存区里还没有可合并的文件（支持 PDF 和图片）")

    writer = PdfWriter()
    pages = 0
    tmpdir = tempfile.mkdtemp(prefix="merge_")
    try:
        for it in picked:
            src = os.path.join(MERGE_DIR, it["id"])
            if it["kind"] == "pdf":
                reader = PdfReader(src)
                if getattr(reader, "is_encrypted", False):
                    try:
                        reader.decrypt("")
                    except Exception:
                        pass
                for page in reader.pages:
                    writer.add_page(page)
                pages += len(reader.pages)
                continue

            # 图片：先画到 A4 白纸上，存成临时单页 PDF，再把这页搬过来
            canvas, dpi = render_image_page(src)
            tmp_pdf = os.path.join(tmpdir, "img_%03d.pdf" % pages)
            canvas.save(tmp_pdf, "PDF", resolution=dpi)
            canvas.close()
            one = PdfReader(tmp_pdf)
            for page in one.pages:
                writer.add_page(page)
            pages += len(one.pages)

        if pages == 0:
            raise ValueError("没拼出任何页面，请检查添加的文件")

        out_name = safe_fname(out_name or "证明材料")
        if not out_name.lower().endswith(".pdf"):
            out_name += ".pdf"
        target = os.path.join(ATTACH_DIR, out_name)
        base, ext = os.path.splitext(out_name)
        i = 1
        while os.path.exists(target):
            target = os.path.join(ATTACH_DIR, "%s(%d)%s" % (base, i, ext))
            i += 1

        tmp_out = target + ".part"
        with open(tmp_out, "wb") as f:
            writer.write(f)
        os.replace(tmp_out, target)
        return os.path.basename(target), pages, os.path.getsize(target)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def norm_atts(v):
    """把附件名单统一成文件名列表。

    以前是把传进来的东西直接当列表迭代 —— 传个字符串会被按字拆开
    （"不是列表" → ['不','是','列','表']），这条从此发不出去，还报一串看不懂的文件名。
    """
    if v is None:
        return []
    if isinstance(v, str):
        v = re.split(r"[,，、;\n|]+", v)
    if not isinstance(v, (list, tuple, set)):
        return []
    out = []
    for x in v:
        s = str(x).strip()
        if s and s not in out:
            out.append(s)
    return out


def norm_email(v):
    """收件邮箱去掉所有空白字符。

    Excel 单元格里粘进来的邮箱可能带换行（Alt+Enter），带着换行发信会被邮件库
    直接判成"头部里藏了别的头"，用户只会看到一串看不懂的英文报错。
    """
    return re.sub(r"\s+", "", str(v or ""))


# 这些字符不可能出现在一个正常邮箱里，出现就是粘错了东西
EMAIL_BAD_CHARS = ':",;<>[]()\\'


def email_problem(v):
    """检查收件邮箱能不能用，没问题返回 ""，有问题返回一句人话。

    单独抽出来是因为"格式不对"以前只在界面上用一句 "要有 @" 糊过去，
    从 Excel 导进来的怪邮箱则一路留到发信时才炸，报的还是英文底层错误。
    """
    e = norm_email(v)
    if not e:
        return "收件邮箱是空的"
    if "@" not in e or e.startswith("@") or e.endswith("@"):
        return "收件邮箱格式不对（%s）" % e[:40]
    bad = sorted({c for c in e if c in EMAIL_BAD_CHARS})
    if bad:
        return "收件邮箱里有不该出现的字符「%s」（%s）" % ("".join(bad), e[:40])
    return ""


def resolve_attachments(job):
    """按投递目标解析附件：atts 为空 -> 发全部附件；否则只发勾选的（按勾选顺序）。"""
    chosen = [str(x) for x in ((job or {}).get("atts") or [])]
    all_atts = list_attachments()
    if not chosen:
        return all_atts
    keep = set(chosen)
    picked = [a for a in all_atts if a["name"] in keep]
    order = {n: i for i, n in enumerate(chosen)}
    picked.sort(key=lambda a: order.get(a["name"], 999))
    return picked


def attachment_problem(job):
    """发送前先检查这一条的附件到底能不能发出去。没问题返回 ""，有问题返回原因。

    为什么非要拦：以前附件对不上时，邮件会"没带简历照样发出去"，状态还记成「已发送」，
    用户根本看不出来。现在宁可这一条不发、明确报错。
    """
    chosen = [str(x) for x in ((job or {}).get("atts") or [])]
    have = {a["name"] for a in list_attachments()}
    if not chosen:
        if not have:
            return "附件库里一个文件都没有，请先到「发送设置」上传简历"
        return ""
    missing = [n for n in chosen if n not in have]
    if missing:
        return "指定的附件在附件库里找不到：%s（被删掉或改过名了？）" % "、".join(missing[:3])
    return ""


def build_message(job, cfg, tpls):
    """构建邮件。

    主题和正文从哪来（优先级从高到低）：
      1. 这条投递目标单独定制的（subject_override / body_override）
      2. 这条「投递定位」指定用哪套话术
      3. 模板库里的默认那套
    """
    tpl = template_for_job(job, tpls)
    subject = render(job.get("subject_override") or tpl.get("subject", ""), job, cfg)
    body = render(job.get("body_override") or tpl.get("body", ""), job, cfg)
    display = cfg.get("display_name") or cfg.get("name") or "求职者"

    msg = MIMEMultipart()
    msg["From"] = formataddr((str(Header(display, "utf-8")), cfg.get("email", "")))
    msg["To"] = str(job.get("email") or "")
    msg["Subject"] = Header(subject, "utf-8")
    msg.attach(MIMEText(body, "plain", "utf-8"))

    for att in resolve_attachments(job):
        with open(os.path.join(ATTACH_DIR, att["name"]), "rb") as f:
            data = f.read()
        part = MIMEApplication(data)
        part.add_header("Content-Disposition", "attachment",
                        filename=("utf-8", "", att["name"]))
        msg.attach(part)
    return msg, subject, tpl.get("name", "")


def send_one(job, cfg, tpls):
    bad = email_problem(job.get("email"))
    if bad:
        return False, "这一条发不了：%s，先在清单里改好" % bad, "", ""
    try:
        msg, subject, tpl_name = build_message(job, cfg, tpls)
        # 90 秒：附件大的时候上传本身就要几十秒。超时太短的话，
        # 邮件可能已经送出去了、这边却记成"失败"，你再重投就会发两遍。
        with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=90) as s:
            s.login(cfg.get("email", ""), cfg.get("auth_code", ""))
            s.send_message(msg)
        return True, "", subject, tpl_name
    except smtplib.SMTPAuthenticationError:
        return False, "登录失败：邮箱地址或授权码不正确（授权码不是QQ密码，需在QQ邮箱设置里生成）", "", ""
    except (smtplib.SMTPException, socket.timeout, OSError) as e:
        return False, "发送失败：%s" % e, "", ""
    except Exception as e:
        # 兜底：任何没预想到的毛病都算这一封失败，不能让整批静默中断
        return False, "发送失败：%s: %s" % (type(e).__name__, e), "", ""


# ---------------------------------------------------------------- 发送线程
SEND_STATE = {
    "running": False,
    "stop": False,
    "current": "",
    "done": 0,
    "total": 0,
    "log": [],  # {time, text, level}
}
# 启动投递时用它把"检查是否在跑"和"标记为在跑"锁成一步。
# 不然两次请求几乎同时进来，都读到"没在跑"，就会同时起两个发送任务 —— 同一批邮件发两遍。
SEND_LOCK = threading.Lock()


def log_add(text, level="info"):
    SEND_STATE["log"].append({
        "time": datetime.now().strftime("%H:%M:%S"),
        "text": text,
        "level": level,
    })
    if len(SEND_STATE["log"]) > 200:
        SEND_STATE["log"] = SEND_STATE["log"][-200:]


def _interruptible_sleep(seconds):
    for _ in range(int(seconds * 2)):
        if SEND_STATE["stop"]:
            return
        time.sleep(0.5)


def int_setting(cfg, key, default):
    """读一个"应该是数字"的设置。配置里被手改成了空字符串/文字也不会炸。"""
    try:
        return int(cfg.get(key, default))
    except (TypeError, ValueError):
        return default


def sender_worker(job_ids):
    """后台一封封发。

    注意：这里**不再**一次性把整份清单读进内存然后整份写回。
    每发一封都重新从磁盘读这一条（用户刚删掉的就不会再发），
    只改这一条的几个字段再写回 —— 这样你在网页上改进展、备注、标签都不会被盖掉。

    整个函数体都包在 try 里：读设置、算间隔这些"开头几步"以前写在 try 外面，
    配置里数字一坏，线程会在还没进循环时就死掉 —— 而 running 已经被置成 True，
    于是界面永远显示"投递中"、再点开始还会被自己拦下，只能重启软件。
    """
    try:
        cfg = get_config()
        tpls = get_templates()
        interval = max(MIN_SEND_INTERVAL, int_setting(cfg, "interval", 40))
        daily_limit = max(1, int_setting(cfg, "daily_limit", 50))
        log_add("开始投递，共 %d 封，每封间隔 %d 秒" % (len(job_ids), interval))
        for idx, jid in enumerate(job_ids):
            if SEND_STATE["stop"]:
                log_add("已手动停止", "err")
                break
            job = find_job(jid)
            if not job:
                log_add("这一条已经被删掉了，跳过", "err")
                SEND_STATE["done"] += 1
                continue
            if today_sent_count() >= daily_limit:
                log_add("已达今日发送上限（%d 封），剩余目标已跳过，明天再继续" % daily_limit, "err")
                break

            # 附件对不上就这一条不发，明确写清原因（绝不发"没带简历"的邮件还记成成功）
            problem = attachment_problem(job)
            if problem:
                err = "未发送：%s" % problem
                mutate_job(jid, lambda j, e=err: j.update({"status": "失败", "error": e}))
                append_sent_log({"id": jid, "company": job.get("company"),
                                 "email": job.get("email"), "ok": False, "error": err,
                                 "sent_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")})
                log_add("%s：%s" % (job.get("company"), err), "err")
                SEND_STATE["done"] += 1
                continue

            SEND_STATE["current"] = "%s · %s" % (job.get("company", ""), job.get("position", ""))
            if not mutate_job(jid, lambda j: j.update({"status": "发送中", "error": ""})):
                log_add("这一条刚好被删掉了，跳过", "err")     # 用户刚删的，就别再发了
                SEND_STATE["done"] += 1
                continue

            ok, err, subject, tpl_name = send_one(job, cfg, tpls)
            now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            def _apply(j, ok=ok, err=err, now=now):
                if ok:
                    j.update({"status": "已发送", "sent_at": now, "error": ""})
                else:
                    j.update({"status": "失败", "error": err})

            mutate_job(jid, _apply)
            append_sent_log({
                "id": jid, "company": job.get("company"), "email": job.get("email"),
                "ok": ok, "error": err, "sent_at": now,
            })
            if ok:
                log_add("已发送：%s → %s（用「%s」｜主题：%s）"
                        % (job.get("company"), job.get("email"), tpl_name, subject), "ok")
            else:
                log_add("%s：%s" % (job.get("company"), err), "err")
            SEND_STATE["done"] += 1

            if idx < len(job_ids) - 1 and not SEND_STATE["stop"]:
                _interruptible_sleep(interval)
    except Exception as e:
        log_add("投递过程中出了意外，已停下来：%s: %s" % (type(e).__name__, e), "err")
    finally:
        SEND_STATE["running"] = False
        SEND_STATE["current"] = ""
        log_add("本轮投递结束：成功处理 %d / %d" % (SEND_STATE["done"], SEND_STATE["total"]))


# ---------------------------------------------------------------- Excel
def excel_template_bytes():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "投递清单"
    headers = ["公司", "岗位", "收件邮箱", "备注"]
    fills = PatternFill("solid", fgColor="185FA5")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fills
    samples = [
        ["示例科技", "后端开发", "hr@example.com", "官网看到的，要求主题写姓名+学校"],
        ["示例工程局", "2027届校园招聘", "join@example2.com", "没写具体岗位就填大类；要求特殊的话回软件里点编辑单独改"],
    ]
    for r, row in enumerate(samples, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    widths = [24, 20, 34, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.cell(row=5, column=1, value="说明：表头请保留这四列（备注可留空），从第2行开始填，填完保存后在本软件中导入。")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_export_bytes():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "投递记录"
    headers = ["公司", "岗位", "收件邮箱", "备注", "投递定位", "标签", "状态", "发送时间",
               "失败原因", "我的进展", "重要日期", "我的笔记"]
    fills = PatternFill("solid", fgColor="185FA5")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fills
    tpls = get_templates()
    for r, j in enumerate(get_jobs(), 2):
        row = [j.get("company"), j.get("position"), j.get("email"), j.get("note"),
               tpl_display_name(j, tpls),
               "，".join(j.get("tags") or []),
               j.get("status"), j.get("sent_at", ""), j.get("error", ""),
               j.get("progress", "未回音"), j.get("event_date", ""), j.get("mynote", "")]
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    widths = [24, 20, 34, 28, 14, 18, 10, 20, 36, 12, 14, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_text(v):
    """把 Excel 单元格的值变成纯文字。日期单独处理，不然会读成 2026-09-16 00:00:00。"""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


# 常见表头写法。用来认「这一行是不是表头」，要求整行每个格子都命中才认，
# 免得把「中铁九局 / 试验检测」这种真实数据误当成表头吃掉。
HEADER_WORDS = {
    "公司", "公司名称", "企业", "企业名称", "单位", "单位名称", "用人单位",
    "岗位", "岗位名称", "职位", "职位名称", "职务", "应聘岗位", "招聘岗位",
    "收件邮箱", "邮箱", "邮箱地址", "电子邮箱", "邮件", "投递邮箱", "email", "mail", "e-mail",
    "备注", "说明", "备注说明", "其他", "序号",
}


def _looks_like_header(vals):
    """判断一行是不是表头：每个格子都是「公司 / 岗位 / 邮箱 / 备注」这类通用词，且都不含 @。"""
    cells = [(_cell_text(v) or "").strip().lower() for v in vals]
    cells = [c for c in cells if c]
    if not cells:
        return False
    if any("@" in c for c in cells):
        return False
    return all(c in HEADER_WORDS for c in cells)


def _append_job(jobs, existing, company, position, email, note, tags,
                progress="", event_date="", mynote="", template_id=""):
    """把一行数据变成清单条目。返回 "added"（新增）或 "dup"（重复，跳过）。"""
    email = norm_email(email)          # Excel 里粘来的邮箱可能带换行，先清干净
    key = (email, position, company)
    if key in existing:
        return "dup"
    existing.add(key)
    jobs.append({
        "id": "j%d%03d" % (int(time.time() * 1000), len(jobs) % 1000),
        "company": company, "position": position, "email": email, "note": note,
        "tags": list(tags or []),
        "template_id": str(template_id or ""),
        "subject_override": "", "body_override": "", "atts": [],
        "progress": progress if progress in ("未回音", "笔试", "面试", "Offer", "挂了", "我放弃") else "未回音",
        "event_date": event_date,
        "mynote": mynote,
        "status": "待发", "error": "", "sent_at": "",
    })
    return "added"


def excel_import(b64_data, tags=None):
    """导入 Excel 文件（一键入口：选完文件直接导）。

    认得出表头（含「公司/岗位/邮箱/备注」）就按表头对列；
    认不出表头也不报错，改成按列顺序硬认 —— 第1列公司、第2列岗位、第3列收件邮箱、第4列备注。
    所以「没写表头、直接从数据开始」的表也能导进来。

    tags 会打给这一批新导入的条目（可为空）。
    """
    from openpyxl import load_workbook
    raw = base64.b64decode(b64_data)
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active

    header_map = {}
    for cell in ws[1]:
        v = _cell_text(cell.value)
        if not v:
            continue
        if "公司" in v or "企业" in v:
            header_map["company"] = cell.column
        elif "岗位" in v or "职位" in v:
            header_map["position"] = cell.column
        elif "邮箱" in v or "mail" in v.lower():
            header_map["email"] = cell.column
        elif "备注" in v or "说明" in v:
            header_map["note"] = cell.column
        elif "进展" in v:
            header_map["progress"] = cell.column
        elif "日期" in v:
            header_map["event_date"] = cell.column
        elif "笔记" in v:
            header_map["mynote"] = cell.column
        elif "定位" in v or "话术" in v or "模板" in v:
            header_map["template"] = cell.column

    if "email" in header_map and ("company" in header_map or "position" in header_map):
        start_row = 2                                  # 表头认出来了，数据从第 2 行开始
    else:
        header_map = {"company": 1, "position": 2, "email": 3, "note": 4}
        start_row = 2 if _looks_like_header([c.value for c in ws[1]]) else 1

    tags, tags_dropped = split_tags(tags)
    tpls = get_templates()
    tpl_by_name = {t["name"]: t["id"] for t in tpls["list"]}
    added, skipped, bad, tpl_miss = 0, 0, 0, 0

    with DATA_LOCK:                                    # 读最新 → 追加 → 写回，整段锁住
        jobs = get_jobs()
        existing = {(j.get("email"), j.get("position"), j.get("company")) for j in jobs}
        for row in ws.iter_rows(min_row=start_row):
            vals = {k: (_cell_text(row[c - 1].value) if 1 <= c <= len(row) else "")
                    for k, c in header_map.items()}
            company = vals.get("company", "")
            position = vals.get("position", "")
            email = vals.get("email", "")
            note = vals.get("note", "")
            if not company and not position and not email and not note:
                continue                               # 整行空着，跳过
            if email_problem(email):
                bad += 1                               # 没有邮箱或格式不对，发不了，算无效
                continue
            # 「投递定位」那一列：写模板名字；写「默认」或留空就跟随默认那套
            tname = (vals.get("template") or "").strip()
            tpl_id = ""
            if tname and tname != DEFAULT_TPL_NAME:
                tpl_id = tpl_by_name.get(tname, "")
                if not tpl_id:
                    tpl_miss += 1                      # 名字对不上，就当默认处理，不挡导入
            r = _append_job(jobs, existing, company, position, email, note, tags,
                            vals.get("progress", ""), vals.get("event_date", ""),
                            vals.get("mynote", ""), tpl_id)
            if r == "dup":
                skipped += 1
            else:
                added += 1
        save_jobs(jobs)
    return {"added": added, "skipped": skipped, "invalid": bad, "tpl_miss": tpl_miss,
            "tags_dropped": tags_dropped}


def import_rows(rows, tags=None):
    """在线填表用：按你框选的那块区域导入。

    列的顺序固定 —— 第1列公司、第2列岗位、第3列收件邮箱、第4列备注。
    框了几列就认几列（少填的列留空），第 5 列往后不看。
    框进去的第一行如果长得像表头（写着「公司/岗位/邮箱」），自动跳过。
    tags 会打给这一批。
    """
    if not isinstance(rows, list) or not rows:
        raise ValueError("没有选中任何内容，先在表格里框一块区域")

    tags, tags_dropped = split_tags(tags)
    added, skipped, bad, header_skip = 0, 0, 0, 0

    with DATA_LOCK:                                    # 同上：整段锁住，别和发信线程互相覆盖
        jobs = get_jobs()
        existing = {(j.get("email"), j.get("position"), j.get("company")) for j in jobs}
        for i, raw in enumerate(rows):
            vals = [(raw[c] if isinstance(raw, list) and c < len(raw) else "") for c in range(len(SHEET_COLS))]
            vals = [_cell_text(v) for v in vals]
            if not header_skip and i < 3 and _looks_like_header(vals):
                header_skip = 1                        # 框进来的表头行，自动跳过（最多跳一行）
                continue
            company, position, email, note = vals
            if not any(vals):
                continue
            if email_problem(email):
                bad += 1
                continue
            r = _append_job(jobs, existing, company, position, email, note, tags)
            if r == "dup":
                skipped += 1
            else:
                added += 1
        save_jobs(jobs)
    return {"added": added, "skipped": skipped, "invalid": bad, "header_skipped": header_skip,
            "tags_dropped": tags_dropped}


# ---------------------------------------------------------------- 在线填表草稿
SHEET_MAX_ROWS = 400
SHEET_MAX_COLS = 8


def sheet_load():
    """读回上次没导完的在线表格内容。"""
    d = load_json(SHEET_FILE, {})
    rows = d.get("rows") if isinstance(d, dict) else None
    if not isinstance(rows, list):
        rows = []
    out = []
    for r in rows[:SHEET_MAX_ROWS]:
        if not isinstance(r, list):
            continue
        out.append([str(c) if c is not None else "" for c in r[:SHEET_MAX_COLS]])
    stamp = d.get("updated_at", "") if isinstance(d, dict) else ""
    return {"rows": out, "updated_at": stamp or ""}


def sheet_save(rows):
    """把在线表格的内容存下来（纯草稿，不影响清单）。"""
    clean = []
    for r in (rows or [])[:SHEET_MAX_ROWS]:
        if not isinstance(r, list):
            continue
        clean.append([str(c)[:500] if c is not None else "" for c in r[:SHEET_MAX_COLS]])
    save_json(SHEET_FILE, {"rows": clean, "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")})
    return clean


def parse_xlsx(b64_data):
    """把 Excel 文件读成一张纯文字的二维表，铺进网页里的在线表格让你看/改。"""
    from openpyxl import load_workbook
    raw = base64.b64decode(b64_data)
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(max_row=SHEET_MAX_ROWS):
        rows.append([_cell_text(c.value) for c in r[:SHEET_MAX_COLS]])
    while rows and not any(rows[-1]):
        rows.pop()                                     # 去掉末尾的空行
    return rows


# ---------------------------------------------------------------- HTTP 服务
class Handler(BaseHTTPRequestHandler):
    server_version = "ResumeSender/1.0"

    def log_message(self, *args):
        pass

    # ---------- 基础 ----------
    def _json(self, obj, status=200):
        data = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _file(self, path, ctype="text/html; charset=utf-8", download_name=None):
        try:
            with open(path, "rb") as f:
                data = f.read()
        except Exception:
            self.send_error(404)
            return
        self.send_response(200)
        self.send_header("Content-Type", ctype)
        if not download_name:
            # 页面本身不缓存：改完代码按一下刷新就能看到新的，不用清缓存
            self.send_header("Cache-Control", "no-store, must-revalidate")
        if download_name:
            self.send_header("Content-Disposition",
                             "attachment; filename*=UTF-8''%s" % quote(download_name))
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _body(self):
        length = int(self.headers.get("Content-Length", 0) or 0)
        raw = self.rfile.read(length) if length else b"{}"
        try:
            return json.loads(raw.decode("utf-8") or "{}")
        except Exception:
            return {}

    def _serve_inline(self, path, name):
        """把文件原样吐给浏览器，不带"下载"标记 —— PDF 就会被浏览器自带的阅读器打开。"""
        try:
            with open(path, "rb") as f:
                data = f.read()
        except Exception:
            self.send_error(404)
            return
        self.send_response(200)
        self.send_header("Content-Type",
                         PREVIEW_TYPES.get(os.path.splitext(name)[1].lower(),
                                           "application/octet-stream"))
        self.send_header("Content-Disposition",
                         "inline; filename*=UTF-8''%s" % quote(name))
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _preview(self, src, key):
        """网页里预览一个文件。src=att 是附件库，src=merge 是材料合并的暂存区。"""
        if src == "merge":
            stored = os.path.basename(str(key or ""))
            p = os.path.join(MERGE_DIR, stored)
            shown = stored.split("__", 1)[1] if "__" in stored else stored
        else:
            shown = os.path.basename(str(key or ""))
            p = os.path.join(ATTACH_DIR, shown)
        if not shown or not os.path.isfile(p):
            self.send_error(404)
            return
        self._serve_inline(p, shown)

    # ---------- GET ----------
    def do_GET(self):
        u = urlparse(self.path)
        path = u.path
        qs = parse_qs(u.query)
        try:
            if path == "/" or path == "/index.html":
                self._file(INDEX_HTML)
            elif path == "/api/attachments/view":
                self._preview((qs.get("src") or ["att"])[0], (qs.get("name") or [""])[0])
            elif path == "/api/logs":
                # 投递记录页：全量发送记录，最新的排最前
                logs = load_json("sent_log.json", [])
                if not isinstance(logs, list):
                    logs = []
                self._json({"ok": True, "logs": list(reversed(logs))})
            elif path == "/api/jobs/history":
                jid = (qs.get("id") or [""])[0]
                logs = [e for e in load_json("sent_log.json", []) if e.get("id") == jid]
                logs.reverse()
                self._json({"ok": True, "logs": logs})
            elif path == "/api/state":
                cfg = get_config()
                masked = dict(cfg)
                masked["auth_code"] = "******" if cfg.get("auth_code") else ""
                self._json({
                    "ok": True,
                    "config": masked,
                    "templates": get_templates(),
                    "max_tags": MAX_TAGS_PER_JOB,
                    "send": {k: SEND_STATE[k] for k in ("running", "stop", "current", "done", "total")},
                    "today_sent": today_sent_count(),
                    "daily_limit": cfg.get("daily_limit", 50),
                    "attachments": list_attachments(),
                    "merge_files": merge_list(),
                    "warnings": list(DATA_WARNINGS),
                })
            elif path == "/api/jobs":
                self._json({"ok": True, "jobs": get_jobs()})
            elif path == "/api/jobs/template":
                self._file_bytes(excel_template_bytes(), "简历投递清单模板.xlsx")
            elif path == "/api/jobs/export":
                self._file_bytes(excel_export_bytes(), "投递记录导出.xlsx")
            elif path == "/api/sheet":
                self._json({"ok": True, **sheet_load(), "all_tags": tag_all()})
            elif path == "/api/status":
                self._json({
                    "ok": True,
                    "running": SEND_STATE["running"],
                    "stop": SEND_STATE["stop"],
                    "current": SEND_STATE["current"],
                    "done": SEND_STATE["done"],
                    "total": SEND_STATE["total"],
                    "log": SEND_STATE["log"][-200:],
                    "jobs_stamp": jobs_stamp(),
                    "today_sent": today_sent_count(),
                })
            else:
                self.send_error(404)
        except Exception as e:
            self._json({"ok": False, "error": str(e)}, 500)

    def _file_bytes(self, data, name):
        self.send_response(200)
        self.send_header("Content-Type",
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition",
                         "attachment; filename*=UTF-8''%s" % quote(name))
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    # ---------- POST ----------
    def do_POST(self):
        path = urlparse(self.path).path
        body = self._body()
        try:
            if path == "/api/config":
                cfg = get_config()
                for k in ("email", "auth_code", "display_name", "name",
                          "school", "major", "interval", "daily_limit"):
                    if k not in body:
                        continue
                    v = body[k]
                    if k in ("interval", "daily_limit"):
                        try:
                            v = int(str(v).strip())
                        except ValueError:
                            continue                   # 填了非数字就保持原样，别把设置写坏
                        # 间隔有下限：填得再小也不让发那么快（太快会被邮箱当成垃圾邮件，账号会被限）
                        cfg[k] = max(MIN_SEND_INTERVAL if k == "interval" else 5, v)
                        continue
                    if k == "auth_code" and str(v) == "******":
                        continue                       # 网页上显示的就是星号，原样发回来不能当成新授权码
                    # 其它字段允许清空（以前空字符串会被忽略，用户以为清掉了，其实后端还是旧值）
                    cfg[k] = str(v).strip()
                save_json("config.json", cfg)
                self._json({"ok": True})               # 不回传授权码，网页要显示就再读 /api/state（那里是星号）
            elif path == "/api/template":
                # 保存某一套话术的内容。不带 id 就存默认那套（老前端也能用）
                tpls = get_templates()
                tid = str(body.get("id") or "")
                t = next((x for x in tpls["list"] if x["id"] == tid), None) if tid else None
                if t is None:
                    t = find_template(tpls, "")
                if "subject" in body:
                    t["subject"] = str(body["subject"])
                if "body" in body:
                    t["body"] = str(body["body"])
                if "name" in body:
                    t["name"] = (str(body["name"]).strip() or t["name"])[:30]
                save_templates(tpls)
                self._json({"ok": True, "templates": tpls})
            elif path == "/api/templates/add":
                # 新建一套话术：主题沿用当前这套的格式（格式一般不变），正文留空等你写
                tpls = get_templates()
                name = (str(body.get("name") or "").strip() or "新话术")[:30]
                taken = {x["name"] for x in tpls["list"]}
                if name in taken:
                    i = 2
                    while ("%s%d" % (name, i)) in taken:
                        i += 1
                    name = "%s%d" % (name, i)
                new = {
                    "id": _new_tpl_id(),
                    "name": name,
                    "subject": str(body.get("subject") or DEFAULT_TEMPLATE["subject"]),
                    "body": "",
                }
                tpls["list"].append(new)
                save_templates(tpls)
                self._json({"ok": True, "templates": tpls, "id": new["id"]})
            elif path == "/api/templates/rename":
                tpls = get_templates()
                tid = str(body.get("id") or "")
                want = str(body.get("name") or "").strip()[:30]
                if not want:
                    self._json({"ok": False, "error": "名字不能空着"})
                    return
                t = next((x for x in tpls["list"] if x["id"] == tid), None)
                if not t:
                    self._json({"ok": False, "error": "没找到这套话术"})
                    return
                if any(x["id"] != tid and x["name"] == want for x in tpls["list"]):
                    self._json({"ok": False, "error": "已经有同名的话术了，换个名字"})
                    return
                t["name"] = want
                save_templates(tpls)
                self._json({"ok": True, "templates": tpls})
            elif path == "/api/templates/default":
                tpls = get_templates()
                tid = str(body.get("id") or "")
                if not any(x["id"] == tid for x in tpls["list"]):
                    self._json({"ok": False, "error": "没找到这套话术"})
                    return
                tpls["default_id"] = tid
                save_templates(tpls)
                self._json({"ok": True, "templates": tpls})
            elif path == "/api/templates/delete":
                tpls = get_templates()
                tid = str(body.get("id") or "")
                if len(tpls["list"]) <= 1:
                    self._json({"ok": False, "error": "至少得留一套话术，不然邮件没内容可发"})
                    return
                if not any(x["id"] == tid for x in tpls["list"]):
                    self._json({"ok": False, "error": "没找到这套话术"})
                    return
                tpls["list"] = [x for x in tpls["list"] if x["id"] != tid]
                if tpls["default_id"] == tid:
                    tpls["default_id"] = tpls["list"][0]["id"]
                save_templates(tpls)
                # 原来挂在它上面的清单，回到「默认」，用户想改再手动改
                n = []

                def _detach(jobs):
                    for j in jobs:
                        if str(j.get("template_id") or "") == tid:
                            j["template_id"] = ""
                            n.append(True)

                mutate_jobs(_detach)
                self._json({"ok": True, "templates": tpls, "reset": len(n)})
            elif path == "/api/preview":
                job_id = body.get("job_id", "")
                jobs = get_jobs()
                job = next((j for j in jobs if j.get("id") == job_id), None)
                if not job:
                    job = jobs[0] if jobs else {"company": "示例科技", "position": "后端开发",
                                                "email": "hr@example.com", "note": "",
                                                "subject_override": "", "body_override": "", "atts": []}
                cfg = get_config()
                tpls = get_templates()
                tpl = template_for_job(job, tpls)
                display = cfg.get("display_name") or cfg.get("name") or "求职者"
                atts = resolve_attachments(job)
                self._json({
                    "ok": True,
                    "from": "%s <%s>" % (display, cfg.get("email", "未配置邮箱")),
                    "to": job.get("email", ""),
                    "subject": render(job.get("subject_override") or tpl.get("subject", ""), job, cfg),
                    "body": render(job.get("body_override") or tpl.get("body", ""), job, cfg),
                    "template_name": tpl.get("name", ""),
                    "attachments": atts,
                    "custom": bool(job.get("subject_override") or job.get("body_override") or job.get("atts")),
                    "job": job,
                })
            elif path == "/api/jobs/add":
                job_tags, tags_dropped = split_tags(body.get("tags"))
                job = {
                    "id": "j%d" % int(time.time() * 1000),
                    "company": str(body.get("company", "")).strip(),
                    "position": str(body.get("position", "")).strip(),
                    "email": norm_email(body.get("email", "")),
                    "note": str(body.get("note", "")).strip(),
                    "tags": job_tags,
                    "template_id": str(body.get("template_id") or "").strip(),
                    "subject_override": str(body.get("subject_override", "")).strip(),
                    "body_override": str(body.get("body_override", "")),
                    "atts": norm_atts(body.get("atts")),
                    "progress": str(body.get("progress", "")).strip() or "未回音",
                    "event_date": str(body.get("event_date", "")).strip(),
                    "mynote": str(body.get("mynote", "")),
                    "status": "待发", "error": "", "sent_at": "",
                }
                bad = email_problem(job["email"])
                if bad:
                    self._json({"ok": False, "error": bad})
                    return
                mutate_jobs(lambda jobs: jobs.append(job))
                self._json({"ok": True, "job": job, "tags_dropped": tags_dropped})
            elif path == "/api/jobs/update":
                jid = body.get("id")
                if "email" in body:
                    bad = email_problem(body.get("email"))
                    if bad:
                        self._json({"ok": False, "error": bad})
                        return
                tags_dropped = 0
                found = []

                def _upd(j):
                    nonlocal tags_dropped
                    found.append(True)
                    # 单行字段去掉首尾空格；正文和笔记保留原样（里面可能有换行和缩进）
                    for k in ("company", "position", "email", "note", "subject_override",
                              "event_date", "template_id"):
                        if k in body:
                            j[k] = norm_email(body[k]) if k == "email" else str(body[k]).strip()
                    for k in ("body_override", "mynote"):
                        if k in body:
                            j[k] = str(body[k])
                    if "progress" in body:
                        v = str(body["progress"]).strip()
                        j["progress"] = v if v in ("未回音", "笔试", "面试", "Offer", "挂了", "我放弃") else "未回音"
                    if "atts" in body:
                        j["atts"] = norm_atts(body.get("atts"))
                    if "tags" in body:
                        new_tags, tags_dropped = split_tags(body.get("tags"))
                        j["tags"] = new_tags

                mutate_job(jid, _upd)
                if not found:
                    self._json({"ok": False, "error": "这一条已经不在了（可能是另一个页面删掉了），刷新一下看看"})
                    return
                self._json({"ok": True, "tags_dropped": tags_dropped})
            elif path == "/api/jobs/delete":
                ids = {str(x) for x in (body.get("ids") or [])}

                def _del(jobs):
                    jobs[:] = [j for j in jobs if str(j.get("id")) not in ids]

                mutate_jobs(_del)
                self._json({"ok": True})
            elif path == "/api/jobs/reset":
                def _reset_all(jobs):
                    for j in jobs:
                        if j.get("status") in ("失败", "已发送"):
                            j["status"] = "待发"
                            j["error"] = ""

                mutate_jobs(_reset_all)
                self._json({"ok": True})
            elif path == "/api/jobs/import":
                result = excel_import(body.get("data", ""), body.get("tags"))
                self._json({"ok": True, **result, "all_tags": tag_all()})
            elif path == "/api/jobs/tag":
                # 批量打标签：mode=add 追加 / set 覆盖 / remove 摘掉
                ids = {str(x) for x in (body.get("ids") or [])}
                if not ids:
                    self._json({"ok": False, "error": "先勾选要打标签的条目"})
                    return
                want, tags_dropped = split_tags(body.get("tags"))
                mode = str(body.get("mode", "add"))
                if mode in ("add", "remove") and not want:
                    self._json({"ok": False, "error": "先写要打的标签（多个用逗号隔开）"})
                    return
                hit = []

                def _tag(jobs):
                    for j in jobs:
                        if str(j.get("id")) not in ids:
                            continue
                        cur = [t for t in (j.get("tags") or []) if t]
                        if mode == "set":
                            cur = list(want)
                        elif mode == "remove":
                            drop = set(want)
                            cur = [t for t in cur if t not in drop]
                        else:
                            for t in want:
                                if t not in cur:
                                    cur.append(t)
                            cur = cur[:MAX_TAGS_PER_JOB]
                        j["tags"] = cur
                        hit.append(True)

                mutate_jobs(_tag)
                self._json({"ok": True, "count": len(hit), "all_tags": tag_all(),
                            "tags_dropped": tags_dropped, "max_tags": MAX_TAGS_PER_JOB})
            elif path == "/api/sheet":
                rows = sheet_save(body.get("rows"))
                self._json({"ok": True, "rows": rows})
            elif path == "/api/sheet/parse":
                rows = parse_xlsx(body.get("data", ""))
                self._json({"ok": True, "rows": rows})
            elif path == "/api/sheet/import":
                result = import_rows(body.get("rows") or [], body.get("tags"))
                self._json({"ok": True, **result, "all_tags": tag_all()})
            elif path == "/api/attachments":
                name = os.path.basename(str(body.get("name", "附件")))
                data = base64.b64decode(body.get("data", ""))
                if len(data) > 15 * 1024 * 1024:
                    self._json({"ok": False, "error": "附件超过 15MB，建议只发 PDF 简历"})
                    return
                target = os.path.join(ATTACH_DIR, name)
                i = 1
                base_name, ext = os.path.splitext(name)
                while os.path.exists(target):
                    target = os.path.join(ATTACH_DIR, "%s(%d)%s" % (base_name, i, ext))
                    i += 1
                with open(target, "wb") as f:
                    f.write(data)
                self._json({"ok": True, "attachments": list_attachments()})
            elif path == "/api/attachments/delete":
                name = os.path.basename(str(body.get("name", "")))
                p = os.path.join(ATTACH_DIR, name)
                if os.path.exists(p):
                    os.remove(p)
                drop_att_note(name)
                self._json({"ok": True, "attachments": list_attachments()})
            elif path == "/api/attachments/note":
                name = os.path.basename(str(body.get("name", "")))
                if name and os.path.isfile(os.path.join(ATTACH_DIR, name)):
                    set_att_note(name, str(body.get("note", "")))
                self._json({"ok": True, "attachments": list_attachments()})
            elif path == "/api/merge/add":
                data = base64.b64decode(body.get("data", ""))
                if not data:
                    self._json({"ok": False, "error": "文件是空的，没收到内容"})
                    return
                if len(data) > 30 * 1024 * 1024:
                    self._json({"ok": False, "error": "单个文件超过 30MB，请先压缩一下"})
                    return
                merge_add(body.get("name", ""), data)
                self._json({"ok": True, "files": merge_list()})
            elif path == "/api/merge/move":
                merge_move(str(body.get("id", "")), int(body.get("dir", 0) or 0))
                self._json({"ok": True, "files": merge_list()})
            elif path == "/api/merge/remove":
                merge_remove(str(body.get("id", "")))
                self._json({"ok": True, "files": merge_list()})
            elif path == "/api/merge/clear":
                merge_clear()
                self._json({"ok": True, "files": merge_list()})
            elif path == "/api/merge/open-tool":
                try:
                    open_merge_tool(str(body.get("name", "")))
                except ValueError as e:
                    self._json({"ok": False, "error": str(e)})
                    return
                self._json({"ok": True})
            elif path == "/api/merge/run":
                try:
                    out, pages, size = merge_run(str(body.get("name", "")))
                except ValueError as e:
                    self._json({"ok": False, "error": str(e)})
                    return
                merge_clear()
                self._json({"ok": True, "output": out, "pages": pages, "size": size,
                            "attachments": list_attachments(), "files": merge_list()})
            elif path == "/api/start":
                # 整段加锁：检查"是否在跑"和标记"开始跑"必须是同一口气完成的，
                # 否则连点两下会同时起两个发送任务，同一批邮件发两遍。
                with SEND_LOCK:
                    if SEND_STATE["running"]:
                        self._json({"ok": False, "error": "已有投递任务在进行中，等它跑完再投"})
                        return
                    cfg = get_config()
                    if not cfg.get("email") or not cfg.get("auth_code"):
                        self._json({"ok": False, "error": "请先在「发送设置」中填写 QQ 邮箱和授权码"})
                        return
                    ids = body.get("ids", [])
                    jobs = get_jobs()
                    valid = [j for j in jobs
                             if j.get("id") in set(ids) and j.get("status") != "已发送"]
                    if not valid:
                        self._json({"ok": False, "error": "没有可投递的目标（已发送的不会重复投递）"})
                        return
                    # 这一批要用到的每套话术都得填好主题和正文，否则发出去是空白邮件
                    tpls = get_templates()
                    bad_tpl = []
                    for j in valid:
                        t = template_for_job(j, tpls)
                        if not str(t.get("subject") or "").strip() or not str(t.get("body") or "").strip():
                            nm = t.get("name") or "未命名"
                            if nm not in bad_tpl:
                                bad_tpl.append(nm)
                    if bad_tpl:
                        self._json({"ok": False,
                                    "error": "「%s」这套话术的主题或正文还是空的，先去「邮件模板」写完再投"
                                             % "、".join(bad_tpl)})
                        return
                    # 附件预检：附件对不上的绝不发出去（以前会静默发一封没带简历的邮件）。
                    # 默认先拦下来问一句；用户在网页上选"只投能发的"才带着 skip_bad 再来一次。
                    bad, good = [], []
                    for j in valid:
                        p = attachment_problem(j)
                        (bad if p else good).append((j, p))
                    if bad and not body.get("skip_bad"):
                        lines = ["%s：%s" % (j.get("company") or j.get("email") or "未命名", p)
                                 for j, p in bad[:5]]
                        more = "，还有 %d 条" % (len(bad) - 5) if len(bad) > 5 else ""
                        self._json({
                            "ok": False,
                            "problems": [{"id": j.get("id"),
                                          "name": j.get("company") or j.get("email") or "未命名",
                                          "why": p} for j, p in bad],
                            "good_count": len(good),
                            "error": "有 %d 条目标的附件对不上，先别发：\n\n%s%s\n\n"
                                     "处理办法：到「发送设置」把附件补上，"
                                     "或者点这些公司的「编辑」重新勾选附件。"
                                     % (len(bad), "\n".join(lines), more)})
                        return
                    if bad:                                  # 用户选了"只投能发的"，这几条标成失败并写清原因
                        bad_ids = {j.get("id") for j, _ in bad}
                        reasons = {j.get("id"): p for j, p in bad}

                        def _mark(jobs_all):
                            for j in jobs_all:
                                if j.get("id") in bad_ids:
                                    j["status"] = "失败"
                                    j["error"] = "未发送：%s" % reasons[j["id"]]

                        mutate_jobs(_mark)
                        for j, p in bad:
                            log_add("跳过（附件问题）：%s —— %s"
                                    % (j.get("company") or j.get("email"), p), "err")
                        valid = [j for j, _ in good]
                        if not valid:
                            self._json({"ok": False, "error": "选中的目标附件都有问题，一条都发不了"})
                            return
                    valid_ids = [j["id"] for j in valid]
                    vs = set(valid_ids)

                    def _reset(jobs_all):
                        for j in jobs_all:
                            # 失败的重发；上次卡在「发送中」的（比如中途关了窗口）也放回待发，
                            # 不然那条的状态会永远停在"发送中"
                            if j.get("id") in vs and j.get("status") in ("失败", "发送中"):
                                j["status"] = "待发"
                                j["error"] = ""

                    mutate_jobs(_reset)
                    # 同步把状态置成"在跑"，窗口期消失，再启动线程
                    SEND_STATE["running"] = True
                    SEND_STATE["stop"] = False
                    SEND_STATE["current"] = ""
                    SEND_STATE["done"] = 0
                    SEND_STATE["total"] = len(valid_ids)
                    SEND_STATE["log"] = []          # 上一轮的日志清掉，别混进来
                    threading.Thread(target=sender_worker, args=(valid_ids,), daemon=True).start()
                self._json({"ok": True, "count": len(valid_ids)})
            elif path == "/api/stop":
                SEND_STATE["stop"] = True
                self._json({"ok": True})
            else:
                self.send_error(404)
        except Exception as e:
            self._json({"ok": False, "error": str(e)}, 500)


LOCK_FILE = os.path.join(DATA_DIR, "app.lock")
APP_PORTS = (8765, 8766, 8767, 8768)      # 8765 被占就顺着往后试


def acquire_instance_lock():
    """保证同时只有一个软件在跑。

    为什么需要：以前双击两次 Start.bat 会开两个软件，它们共用同一份 data 目录，
    互相覆盖数据、而且两边都能点"开始投递" —— 同一批邮件发两遍。
    拿不到锁就返回 None（说明已经有一个在跑了）。
    """
    f = open(LOCK_FILE, "a+", encoding="utf-8")
    try:
        if os.name == "nt":
            import msvcrt
            f.seek(0)
            msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(f.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        f.close()
        return None
    return f


def read_running_port():
    """从锁文件里读出那个实例开的端口（Linux/macOS 上能用；Windows 上读不到，见下面）。"""
    try:
        with open(LOCK_FILE, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line.isdigit():
                    return int(line)
    except OSError:
        pass
    return None


def find_running_instance():
    """挨个端口问一句"你是谁"，认出正在跑的本软件。

    为什么不靠锁文件里的端口号：Windows 上锁文件第 0 个字节被 msvcrt 上了锁，
    别的进程去读会直接 PermissionError（同一个进程开第二个句柄都读不出来），
    结果永远回落到 8765 —— 万一 8765 被别的程序占着、主实例退到了 8766，
    再双击 Start.bat 就会把别人家的网页打开。
    """
    import http.client
    for p in APP_PORTS:
        try:
            c = http.client.HTTPConnection("127.0.0.1", p, timeout=1.5)
            c.request("GET", "/api/state")
            r = c.getresponse()
            server = str(r.getheader("Server") or "")
            r.read(64)
            c.close()
        except Exception:
            continue
        if server.startswith("ResumeSender/"):
            return p
    return None


def main():
    os.makedirs(DATA_DIR, exist_ok=True)       # 先建目录，锁文件要放里面
    lock = acquire_instance_lock()
    if lock is None:
        port = find_running_instance() or read_running_port() or APP_PORTS[0]
        url = "http://127.0.0.1:%d" % port
        print("=" * 50)
        print("  软件已经在运行了，不用再开一个。")
        print("  正在给你打开已有的那个窗口：%s" % url)
        print("  如果浏览器没反应，就手动复制上面这个网址打开。")
        print("=" * 50)
        webbrowser.open(url)
        time.sleep(1.5)
        return

    init_dirs()                                # 确认只有自己在跑，才动数据目录
    server = None
    port = APP_PORTS[0]
    for p in APP_PORTS:
        try:
            server = ThreadingHTTPServer(("127.0.0.1", p), Handler)
            port = p
            break
        except OSError:
            continue
    if server is None:
        print("错误：8765-8768 端口都被占用，无法启动")
        return
    try:
        lock.seek(0)
        lock.truncate()
        lock.write("%d\n" % port)
        lock.flush()
    except OSError:
        pass
    url = "http://127.0.0.1:%d" % port
    print("=" * 50)
    print("  简历批量投递助手 已启动")
    print("  请在浏览器访问：%s" % url)
    print("  关闭本窗口即退出软件（数据都保存在本地）")
    print("=" * 50)
    threading.Timer(1.0, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
